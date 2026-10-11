"""对话附件（问答链路）：对象存储复用 services/storage（MinIO 优先，未配置回退本地磁盘），
元数据与解析文本落 chat_attachments。

上传即解析：pdf/word/excel/ppt/txt/md 同步提取纯文本（超长截断），图片标记 image 走多模态视觉；
发消息时 sessions.chat 按 id 列表绑定——user/message 事件只带 AttachmentMeta（回放轻量），
解析文本仅注入 LLM 上下文，不回传前端。下载走鉴权端点而非静态挂载，storage_key 服务端生成。

对象存储路径按上传日期分层：chat-attachments/{YYYYMM}/{DD}/{user_id}/{attachment_id}/{文件名}，
MinIO 控制台按日期前缀浏览、按月清理归档方便；storage_key 入库，旧文件路径不变天然兼容。"""
from __future__ import annotations

import io
import logging
import mimetypes
import re
import uuid
from datetime import datetime
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core import errors
from ..persistence.models import ChatAttachment
from . import storage

logger = logging.getLogger(__name__)

# 上传白名单：pdf/word/excel/ppt/图片/text/md（.doc/.ppt 旧格式可上传但解析降级）
ALLOWED_EXTS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".txt", ".md"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
DEFAULT_MAX_MB = 20
# 注入 LLM 的单附件文本上限（字符）：控制上下文膨胀，超出截断并注明
PARSED_TEXT_CAP = 80_000
# 图片注入视觉通道的大小上限（字节）：超出仅按名称提及，不构造 image_url
IMAGE_VISION_MAX_BYTES = 8 * 1024 * 1024
_NAME_BAD = re.compile(r'[\\/:*?"<>|\r\n\t]+')


def max_mb() -> int:
    from ..core.config import settings

    return settings.chat_attachment_max_mb or DEFAULT_MAX_MB


def extract_text(file_name: str, content: bytes) -> tuple[str, str | None, str | None]:
    """按扩展名提取文本。返回 (parse_status, parsed_text, parse_error)：
    ok=已提取 / image=走视觉 / failed=无法提取（error 说明原因）。"""
    ext = Path(file_name).suffix.lower()
    try:
        if ext in IMAGE_EXTS:
            return "image", None, None
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        elif ext == ".docx":
            import docx

            d = docx.Document(io.BytesIO(content))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append("\t".join(cells))
            text = "\n".join(parts)
        elif ext == ".doc":
            return "failed", None, "旧版 .doc 暂不支持解析，请另存为 .docx 后重新上传"
        elif ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(io.BytesIO(content))
            slides: list[str] = []
            for idx, slide in enumerate(prs.slides, 1):
                texts: list[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        texts.append(shape.text_frame.text.strip())
                    if shape.has_table:
                        for row in shape.table.rows:
                            cells = [c.text.strip() for c in row.cells]
                            if any(cells):
                                texts.append("\t".join(cells))
                if slide.has_notes_slide:
                    note = slide.notes_slide.notes_text_frame.text.strip()
                    if note:
                        texts.append(f"（备注）{note}")
                if texts:
                    slides.append(f"【幻灯片 {idx}】\n" + "\n".join(texts))
            text = "\n\n".join(slides)
        elif ext == ".ppt":
            return "failed", None, "旧版 .ppt 暂不支持解析，请另存为 .pptx 后重新上传"
        elif ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                lines.append(f"【工作表 {ws.title}】")
                for row in ws.iter_rows(values_only=True):
                    vals = ["" if v is None else str(v) for v in row]
                    if any(vals):
                        lines.append("\t".join(vals))
            wb.close()
            text = "\n".join(lines)
        elif ext == ".xls":
            import xlrd

            book = xlrd.open_workbook(file_contents=content)
            lines = []
            for ws in book.sheets():
                lines.append(f"【工作表 {ws.name}】")
                for r in range(ws.nrows):
                    vals = [str(v) for v in ws.row_values(r)]
                    if any(vals):
                        lines.append("\t".join(vals))
            text = "\n".join(lines)
        elif ext in (".txt", ".md"):
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                text = content.decode("gbk", errors="replace")
        else:
            return "failed", None, "不支持的解析类型"
    except Exception as exc:  # 解析失败不阻断上传：入库 failed + 原因，问答链路降级为仅文件名
        logger.warning("附件解析失败 name=%.80s err=%r", file_name, exc)
        return "failed", None, f"解析失败：{exc}"
    text = text.strip()
    if not text:
        return "failed", None, "未提取到文本（可能为扫描件或空文档）"
    if len(text) > PARSED_TEXT_CAP:
        text = text[:PARSED_TEXT_CAP] + "\n\n…（内容超长已截断）"
    return "ok", text, None


def save_upload(db: Session, *, content: bytes, file_name: str, user_id: str) -> ChatAttachment:
    """校验 → 对象存储 → 解析 → 元数据入库（上传即解析，失败不阻断上传）。"""
    if not content:
        raise errors.validation("空文件")
    safe = _NAME_BAD.sub("_", file_name).strip() or "未命名"
    ext = Path(safe).suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise errors.validation(f"不支持的文件类型：{ext or '（无扩展名）'}")
    limit = max_mb() * 1024 * 1024
    if len(content) > limit:
        raise errors.validation(f"文件超过 {max_mb()}MB 上限")
    mime = mimetypes.guess_type(safe)[0] or "application/octet-stream"
    attachment_id = str(uuid.uuid4())
    # 键含 attachment_id 天然唯一；日期分目录便于 MinIO 控制台浏览与按月归档
    now = datetime.now()
    storage_key = f"chat-attachments/{now:%Y%m}/{now:%d}/{user_id}/{attachment_id}/{safe}"
    storage.put(storage_key, content, content_type=mime)
    status, text, error = extract_text(safe, content)
    row = ChatAttachment(
        attachment_id=attachment_id, user_id=user_id, file_name=safe, mime=mime,
        size=len(content), storage_key=storage_key, parse_status=status,
        parse_error=error, parsed_text=text,
    )
    db.add(row)
    db.flush()
    logger.info("附件上传 user=%s id=%s name=%.80s size=%d status=%s",
                user_id, row.attachment_id, safe, len(content), status)
    return row


def attachment_view(a: ChatAttachment) -> dict:
    """契约 AttachmentMeta 视图（不含 parsed_text，回放/前端零冗余）。"""
    return {"id": a.attachment_id, "name": a.file_name, "mime": a.mime, "size": a.size, "status": a.parse_status}


def get_owned(db: Session, attachment_id: str, user_id: str) -> ChatAttachment:
    """读取本人附件：不存在或不归属一律 404（不泄露存在性）。"""
    row = db.get(ChatAttachment, attachment_id)
    if row is None or row.user_id != user_id:
        raise errors.not_found("附件不存在")
    return row


def resolve_for_chat(db: Session, ids: list[str], user_id: str) -> list[ChatAttachment]:
    """发消息绑定：按给定顺序取本人附件，未知/越权 id 直接拒绝。"""
    if not ids:
        return []
    if len(ids) > 10:
        raise errors.validation("单条消息最多携带 10 个附件")
    if len(set(ids)) != len(ids):
        raise errors.validation("附件重复引用")
    rows = db.scalars(
        select(ChatAttachment).where(ChatAttachment.attachment_id.in_(ids))  # noqa: E711
    ).all()
    by_id = {r.attachment_id: r for r in rows if r.user_id == user_id}
    missing = [i for i in ids if i not in by_id]
    if missing:
        raise errors.validation(f"附件不可用：{','.join(missing[:3])}")
    return [by_id[i] for i in ids]


def read_bytes(a: ChatAttachment) -> bytes:
    return storage.read(a.storage_key)
