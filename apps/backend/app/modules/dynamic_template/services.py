"""动态填报模板服务：Excel 解析 → 字段推断 → HTML 双缓存（填报页/汇总页）→ 派发下拉动态项。

范围约束：读取首个 Sheet；自动识别标题/说明行和真实表头；不做公式重算
（data_only 取缓存值）；类型推断仅是初始值，用户在预览步可调整后提交最终 columns。
"""
from __future__ import annotations

import logging
import math
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from jinja2 import Environment, FileSystemLoader
from openpyxl import load_workbook
from sqlalchemy import or_, select
from sqlalchemy.orm import Session
import xlrd

from ...core import errors
from ...persistence.models import BizTask, BizTodo
from ...services import todo as todo_svc
from .models import DynamicFillData, DynamicTemplate

logger = logging.getLogger(__name__)

# Jinja2 模板目录（本模块 html_templates/，随代码发布）
_TPL_DIR = Path(__file__).resolve().parent / "html_templates"
_env = Environment(loader=FileSystemLoader(str(_TPL_DIR)), autoescape=True)

MAX_COLUMNS = 40          # 单模板字段上限
MAX_SAMPLE_ROWS = 200     # 类型推断采样行数
SELECT_MAX_OPTIONS = 8    # select 推断：去重值上限
SELECT_RATIO = 0.5        # select 推断：去重/非空 比例阈值
TEXTAREA_AVG_LEN = 30     # textarea 推断：平均字符长度阈值

_VALID_TYPES = ("text", "textarea", "number", "date", "select", "table")


# ---------------- Excel 解析与类型推断 ----------------


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _col_key(header: str, idx: int, used: set[str]) -> str:
    """表头 → data_json 的 key：清洗分隔符，冲突加后缀，空表头退化为 col_N。"""
    k = re.sub(r"[\s/\\:：,，。；;、（）()\[\]【】{}'\"?？!！*]+", "_", header).strip("_")
    if not k:
        k = f"col_{idx}"
    base, n = k, 2
    while k in used:
        k, n = f"{base}_{n}", n + 1
    used.add(k)
    return k


def _infer_type(values: list, label: str = "") -> tuple[str, list[str]]:
    """按样本列值推断字段类型；返回 (type, select_options)。"""
    non_empty = [v for v in values if v is not None and str(v).strip() != ""]
    if not non_empty:
        return _infer_type_from_label(label), []
    if all(isinstance(v, datetime) for v in non_empty):
        return "date", []
    if all(_is_number(v) for v in non_empty):
        return "number", []
    texts = [_cell_text(v) for v in non_empty]
    uniq = sorted(set(texts))
    if (
        len(non_empty) >= 4
        and len(uniq) <= SELECT_MAX_OPTIONS
        and len(uniq) / len(non_empty) <= SELECT_RATIO
    ):
        return "select", uniq
    if sum(len(t) for t in texts) / len(texts) > TEXTAREA_AVG_LEN:
        return "textarea", []
    return "text", []


def _infer_type_from_label(label: str) -> str:
    """无样本值时按字段语义保守推断；不自动生成无法确认的下拉选项。"""
    normalized = re.sub(r"\s+", "", label).lower()
    if any(word in normalized for word in ("日期", "时间", "年月", "期限")):
        return "date"
    if any(word in normalized for word in (
        "金额", "投资", "投入", "经费", "支出", "收入", "成本", "利润", "数量", "比例", "比率", "百分比",
    )):
        return "number"
    if any(word in normalized for word in (
        "内容", "目标", "说明", "描述", "情况", "原因", "措施", "意见", "备注", "计划",
    )):
        return "textarea"
    return "text"


def _is_number(v) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    s = str(v).strip().replace(",", "")
    if not s:
        return False
    try:
        float(s)
        return True
    except ValueError:
        return False


def _read_xls(content: bytes) -> tuple[str, list[list], list[float]]:
    """读取 Excel 97-2003 工作簿，并把日期单元格转换为 datetime。"""
    wb = xlrd.open_workbook(file_contents=content, formatting_info=True)
    if wb.nsheets == 0:
        return "", [], []
    ws = wb.sheet_by_index(0)
    rows: list[list] = []
    for row_idx in range(ws.nrows):
        row = []
        for col_idx in range(ws.ncols):
            cell = ws.cell(row_idx, col_idx)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(value, wb.datemode)
            elif cell.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                value = int(value)
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                value = None
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                value = bool(value)
            row.append(value)
        rows.append(row)
    widths = [
        min(420, max(80, round((ws.colinfo_map.get(i).width if i in ws.colinfo_map else 2560) / 32)))
        for i in range(ws.ncols)
    ]
    return ws.name, rows, widths


def _read_modern_excel(content: bytes) -> tuple[str, list[list], list[float]]:
    # 普通模式才能读取列宽等版式信息；上传大小已由 API 限制为 10MB。
    wb = load_workbook(BytesIO(content), read_only=False, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        widths = [
            min(420, max(80, round((ws.column_dimensions[cell.column_letter].width or 10) * 8)))
            for cell in ws[1]
        ] if rows else []
        return ws.title, rows, widths
    finally:
        wb.close()


def parse_excel(content: bytes, file_name: str | None = None) -> dict:
    """解析 Excel → {sheet_name, columns, sample_rows}（解析预览，不落库）。"""
    try:
        if (file_name or "").lower().endswith(".xls"):
            sheet_name, rows, widths = _read_xls(content)
        else:
            sheet_name, rows, widths = _read_modern_excel(content)
    except Exception as exc:
        raise errors.validation(f"Excel 文件无法解析：{exc}") from exc
    if not rows:
        raise errors.validation("Excel 首个 Sheet 为空")
    header_idx, encoded = _detect_header_row(rows)
    headers = [(_cell_text(h) if h is not None else "") for h in rows[header_idx][:MAX_COLUMNS]]
    body = [
        r[:MAX_COLUMNS]
        for r in rows[header_idx + 1:]
        if any(c is not None and str(c).strip() for c in r)
    ]
    if not any(headers):
        raise errors.validation("首行表头为空：请确认第一行是字段名")
    sample = body[:MAX_SAMPLE_ROWS]
    used: set[str] = set()
    columns = []
    for i, h in enumerate(headers):
        if not h:
            continue
        label, required = _parse_header(h) if encoded else (h, i == 0)
        ftype, options = _infer_type([r[i] if i < len(r) else None for r in sample], label)
        columns.append({
            "key": _col_key(label, i, used),
            "label": label,
            "type": ftype,
            "required": required,
            "options": options,
            "source_col": i,
            "width": widths[i] if i < len(widths) else 140,
        })
    if not columns:
        raise errors.validation("首行表头为空：请确认第一行是字段名")
    preamble = [_cell_text(v) for row in rows[:header_idx] for v in row if _cell_text(v)]
    title = _extract_title(preamble[0]) if preamble else sheet_name
    return {
        "sheet_name": sheet_name,
        "columns": columns,
        "sample_rows": [
            {
                c["key"]: _cell_text(r[c["source_col"]])
                if c["source_col"] < len(r) else ""
                for c in columns
            }
            for r in sample[:5]
        ],
        "presentation": "grid" if encoded or len(columns) >= 6 else "form",
        "title": title,
        "metadata": preamble[1:] if preamble else [],
        "header_row": header_idx,
    }


_ENCODED_HEADER = re.compile(r"^\s*[^<>]+<(.+?)>\s*(\(\*\)|（\*）)?\s*$")


def _detect_header_row(rows: list[list]) -> tuple[int, bool]:
    """优先识别 CODE<中文名>(*) 导出格式，否则选择前十行中最像表头的一行。"""
    candidates: list[tuple[int, int, int]] = []
    for idx, row in enumerate(rows[:10]):
        texts = [_cell_text(v) for v in row[:MAX_COLUMNS] if _cell_text(v)]
        encoded_count = sum(bool(_ENCODED_HEADER.match(v)) for v in texts)
        if encoded_count >= 2 and encoded_count >= len(texts) * 0.6:
            return idx, True
        text_cells = sum(
            isinstance(v, str) and bool(v.strip())
            for v in row[:MAX_COLUMNS]
        )
        candidates.append((len(texts), text_cells, idx))
    viable = [c for c in candidates if c[0] >= 2]
    # 列数相同时优先文本占比高且更靠前的行，避免把普通表格的数据行误判成表头。
    return (max(viable, key=lambda c: (c[0], c[1], -c[2]))[2] if viable else 0), False


def _parse_header(value: str) -> tuple[str, bool]:
    match = _ENCODED_HEADER.match(value)
    return (match.group(1).strip(), bool(match.group(2))) if match else (value, False)


def _extract_title(value: str) -> str:
    return value.split(":", 1)[-1].strip() if ":" in value else value.strip()


# ---------------- HTML 渲染 ----------------


def render_fill_html(
    name: str, columns: list[dict], initial: dict | None = None,
    layout: dict | None = None,
) -> str:
    """填报页 HTML：iframe srcdoc 内嵌，postMessage 与父页面通信（submit/cancel/ready）。"""
    return _env.get_template("fill_form.j2").render(
        name=name, columns=columns, initial=initial or {}, layout=layout or {},
    )


def render_summary_html(name: str, columns: list[dict], layout: dict | None = None) -> str:
    """汇总预览 HTML：上传确认页展示字段清单（只读）。"""
    return _env.get_template("summary.j2").render(name=name, columns=columns, layout=layout or {})


# ---------------- 模板 CRUD ----------------


def _validate_columns(columns: list[dict], limit: int = MAX_COLUMNS) -> None:
    if not columns:
        raise errors.validation("字段列表不能为空")
    if len(columns) > limit:
        raise errors.validation(f"字段数量超过上限 {limit}")
    seen: set[str] = set()
    for c in columns:
        k, label, ftype = c.get("key"), c.get("label"), c.get("type")
        if not isinstance(k, str) or not k or not isinstance(label, str) or not label.strip():
            raise errors.validation("字段缺少 key/label")
        if not isinstance(ftype, str) or ftype not in _VALID_TYPES:
            raise errors.validation(f"字段「{label}」类型非法：{ftype}")
        if k in seen:
            raise errors.validation(f"字段 key 重复：{k}")
        seen.add(k)
        if ftype == "select" and not c.get("options"):
            raise errors.validation(f"字段「{label}」为下拉选择但缺少选项")
        if not isinstance(c.get("required"), bool):
            raise errors.validation(f"字段「{label}」的必填设置不合法")
        options = c.get("options", [])
        if not isinstance(options, list) or not all(isinstance(value, str) for value in options):
            raise errors.validation(f"字段「{label}」的选项不合法")


def _view(t: DynamicTemplate, *, with_html: bool = False) -> dict:
    out = {
        "template_id": t.template_id,
        "name": t.name,
        "source_file_name": t.source_file_name,
        "visibility": t.visibility,
        "owner_id": t.owner_id,
        "status": t.status,
        "columns": (t.parse_result_json or {}).get("columns", []),
        "tabs": (t.parse_result_json or {}).get("tabs", []),
        "created_at": t.created_at.isoformat(),
    }
    if with_html:
        out["form_html"] = t.form_html
        out["summary_html"] = t.summary_html
    return out


def create_template(
    db: Session, *, name: str, columns: list[dict], parse_result: dict,
    source_file_name: str | None, owner_id: str,
) -> dict:
    """创建动态模板：渲染两份 HTML 一并落库（创建即启用，出现在派发下拉）。"""
    tabs = parse_result.get("tabs") or []
    _validate_columns(columns, MAX_COLUMNS * len(tabs) if tabs else MAX_COLUMNS)
    if tabs:
        keys = [key for tab in tabs for key in tab["field_keys"]]
        if len(keys) != len(set(keys)) or set(keys) != {c["key"] for c in columns}:
            raise errors.validation("分页字段必须与模板字段一一对应")
        if len(tabs) > 10 or len({tab["id"] for tab in tabs}) != len(tabs):
            raise errors.validation("分页数量超限或标识重复")
        for tab in tabs:
            if not tab["name"].strip():
                raise errors.validation("分页名称不能为空")
            _validate_columns([c for c in columns if c["key"] in tab["field_keys"]])
    if not name.strip():
        raise errors.validation("模板名称不能为空")
    pr = dict(parse_result)
    pr["columns"] = columns
    layout = {
        "presentation": pr.get("presentation", "form"),
        "title": pr.get("title", name),
        "metadata": pr.get("metadata", []),
    }
    row = DynamicTemplate(
        name=name.strip(),
        source_file_name=source_file_name,
        parse_result_json=pr,
        form_html=render_fill_html(name, columns, layout=layout),
        summary_html=render_summary_html(name, columns, layout=layout),
        visibility="system",
        owner_id=owner_id,
        status="active",
    )
    db.add(row)
    db.flush()
    return _view(row, with_html=True)


def get_template_row(db: Session, template_id: str) -> DynamicTemplate:
    row = db.get(DynamicTemplate, template_id)
    if row is None or row.is_delete:
        raise errors.not_found("动态模板不存在")
    return row


def get_template_view(db: Session, template_id: str, *, with_html: bool = False) -> dict:
    return _view(get_template_row(db, template_id), with_html=with_html)


def list_templates(db: Session, user_id: str) -> list[dict]:
    """管理列表：系统级共享 + 本人私有，含 draft/disabled（派发下拉另行只取 active）。"""
    rows = db.scalars(
        select(DynamicTemplate)
        .where(
            DynamicTemplate.is_delete.is_(False),
            or_(DynamicTemplate.visibility == "system", DynamicTemplate.owner_id == user_id),
        )
        .order_by(DynamicTemplate.created_at.desc())
    ).all()
    return [_view(t) for t in rows]


def list_active(db: Session) -> list[dict]:
    """启用中的模板视图（merge_dynamic_into_list 派发下拉数据源）。"""
    rows = db.scalars(
        select(DynamicTemplate)
        .where(
            DynamicTemplate.is_delete.is_(False),
            DynamicTemplate.status == "active",
        )
        .order_by(DynamicTemplate.created_at.desc())
    ).all()
    return [_view(t) for t in rows]


def patch_status(db: Session, template_id: str, status: str) -> dict:
    row = get_template_row(db, template_id)
    if status not in ("active", "disabled"):
        raise errors.validation("状态仅允许 active/disabled")
    row.status = status
    db.flush()
    return _view(row)


def logic_delete(db: Session, template_id: str) -> dict:
    row = get_template_row(db, template_id)
    row.is_delete = True
    db.flush()
    return {"template_id": template_id, "deleted": True}


def merge_dynamic_into_list(db: Session, items: list[dict]) -> list[dict]:
    """静态任务模板视图 + 动态模板 → 派发下拉统一列表（动态项 template_key = dyn_<id>）。"""
    out = list(items)
    for t in list_active(db):
        out.append({
            "template_key": f"dyn_{t['template_id']}",
            "name": t["name"],
            "scene": "dynamic_fill",
            "review_required": True,  # 动态填报默认两段式（财务收集数据需审核）
            "related_system": None,
            "period_rule": "natural_month",
            "creator_roles": [],
            "dynamic": True,
            "template_id": t["template_id"],
        })
    return out


# ---------------- 填报提交 ----------------


def _validate_data(columns: list[dict], data: dict) -> None:
    for c in columns:
        raw = data.get(c["key"])
        values = raw if isinstance(raw, list) else [raw]
        if c.get("required") and (not values or any(_is_empty(v) for v in values)):
            raise errors.validation(f"「{c['label']}」每条记录均为必填项")
        for v in values:
            if _is_empty(v):
                continue
            if c["type"] == "select" and v not in (c.get("options") or []):
                raise errors.validation(f"「{c['label']}」取值不合法")
            if c["type"] == "number" and (isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v)):
                raise errors.validation(f"「{c['label']}」需为数字")


def _validate_row_shape(columns: list[dict], data: dict) -> None:
    values = [data.get(c["key"]) for c in columns]
    if not any(isinstance(value, list) for value in values):
        return
    if not all(isinstance(value, list) for value in values):
        raise errors.validation("同一分页的多行字段格式不一致")
    lengths = {len(value) for value in values}
    if len(lengths) != 1 or lengths == {0}:
        raise errors.validation("同一分页的多行记录数量必须一致且不能为空")


def render_fill_page(db: Session, template_id: str, todo_id: str | None, user_id: str) -> dict:
    """填报页 HTML（返回 JSON，前端 srcdoc 内嵌）：todo_id 已有数据 → 预填（退回重填）。"""
    tpl = get_template_row(db, template_id)
    columns = (tpl.parse_result_json or {}).get("columns", [])
    initial: dict = {}
    if todo_id:
        t = db.get(BizTodo, todo_id)
        if t is None:
            raise errors.not_found("待办不存在")
        if t.assignee_id != user_id:
            raise errors.forbidden("仅承办人本人可填报")
        task = db.get(BizTask, t.task_id)
        if task is None or task.dynamic_template_id != template_id:
            raise errors.validation("填报模板与待办不匹配")
        draft = db.scalars(
            select(DynamicFillData).where(
                DynamicFillData.todo_id == todo_id,
                DynamicFillData.status == "draft",
                DynamicFillData.is_delete.is_(False),
            )
        ).first()
        initial = (draft.data_json or {}) if draft else (t.payload or {})
    # 每次按最新 Jinja2 模板渲染，历史模板无需重新上传即可获得样式与交互升级。
    html = render_fill_html(tpl.name, columns, initial, tpl.parse_result_json or {})
    return {
        "template_id": tpl.template_id, "name": tpl.name, "html": html,
        "tabs": (tpl.parse_result_json or {}).get("tabs", []),
        "columns": columns, "initial": initial,
    }


def save_fill_draft(
    db: Session, *, todo_id: str, template_id: str, data: dict, user_id: str,
) -> dict:
    """保存承办人草稿，不校验必填项，也不推进待办状态。"""
    tpl = get_template_row(db, template_id)
    t = db.get(BizTodo, todo_id)
    if t is None:
        raise errors.not_found("待办不存在")
    if t.assignee_id != user_id:
        raise errors.forbidden("仅承办人本人可保存草稿")
    task = db.get(BizTask, t.task_id)
    if task is None or task.dynamic_template_id != template_id:
        raise errors.validation("填报模板与待办不匹配")
    if t.status not in ("pending", "returned", "submitted"):
        raise errors.validation(f"当前状态 {t.status} 不允许保存草稿")
    allowed_keys = {column["key"] for column in (tpl.parse_result_json or {}).get("columns", [])}
    if set(data) - allowed_keys:
        raise errors.validation("草稿包含模板之外的字段")
    now = datetime.now(timezone.utc)
    row = db.scalars(select(DynamicFillData).where(DynamicFillData.todo_id == todo_id)).first()
    if row is None:
        row = DynamicFillData(todo_id=todo_id, task_id=t.task_id, template_id=template_id)
        db.add(row)
    row.data_json = data
    row.submitted_by = user_id
    row.submitted_at = now
    row.status = "draft"
    db.flush()
    return {"todo_id": todo_id, "status": "draft", "saved_at": now.isoformat()}


def ensure_fill_submitted(
    db: Session, *, todo_id: str, template_id: str, data: dict, user_id: str,
) -> dict:
    """动态填报提交：审核前允许反复修改；审核通过后由状态机锁定。"""
    from ...services import todo as todo_svc

    tpl = get_template_row(db, template_id)
    columns = (tpl.parse_result_json or {}).get("columns", [])
    t = db.get(BizTodo, todo_id)
    if t is None:
        raise errors.not_found("待办不存在")
    if t.assignee_id != user_id:
        raise errors.forbidden("仅承办人本人可提交")
    task = db.get(BizTask, t.task_id)
    if task is None or task.dynamic_template_id != template_id:
        raise errors.validation("填报模板与待办不匹配")
    if t.status not in ("pending", "returned", "submitted"):
        raise errors.validation(f"当前状态 {t.status} 不允许提交")
    tabs = (tpl.parse_result_json or {}).get("tabs", [])
    if tabs:
        for tab in tabs:
            tab_columns = [{**c, "label": f"{tab['name']} / {c['label']}"}
                           for c in columns if c["key"] in tab["field_keys"]]
            _validate_row_shape(tab_columns, data)
            _validate_data(tab_columns, data)
    else:
        _validate_data(columns, data)
    row = db.scalars(
        select(DynamicFillData).where(DynamicFillData.todo_id == todo_id)
    ).first()
    if row is None:
        row = DynamicFillData(todo_id=todo_id, task_id=t.task_id, template_id=template_id)
        db.add(row)
    # submitted 状态重提时覆盖同一条记录，不新增待办或填报记录。
    row.data_json = data
    row.submitted_by = user_id
    row.submitted_at = datetime.now(timezone.utc)
    row.status = "submitted"
    db.flush()
    # payload 同步写 BizTodo（审核弹窗走现有流程即可看到填报内容）
    if t.status in ("pending", "returned"):
        result = todo_svc.submit(db, todo_id, user_id, payload=data)
        db.flush()
        return result
    # 待审核状态仅更新最新数据，保持 submitted，派发人看到的始终是最新版本。
    t.payload = data
    t.lamp_agg = todo_svc._count_lamps(data)
    t.submitted_at = datetime.now(timezone.utc)
    db.flush()
    todo_svc._notify_todo(db, t)
    return todo_svc.todo_view(db, t)


def task_analytics(db: Session, task_id: str, user_id: str) -> dict:
    """派发人数据分析：汇总填写记录，并按模板字段计算填写率及基础分布。"""
    task = db.get(BizTask, task_id)
    if task is None:
        raise errors.not_found("任务不存在")
    if task.dispatcher_id != user_id:
        raise errors.forbidden("仅任务派发人可查看数据分析")
    if not task.dynamic_template_id:
        raise errors.validation("该任务不是动态信息填报任务")

    tpl = get_template_row(db, task.dynamic_template_id)
    columns = (tpl.parse_result_json or {}).get("columns", [])
    todos = db.scalars(select(BizTodo).where(BizTodo.task_id == task_id)).all()
    fills = db.scalars(
        select(DynamicFillData).where(
            DynamicFillData.task_id == task_id,
            DynamicFillData.status == "submitted",
            DynamicFillData.is_delete.is_(False),
        )
    ).all()
    fill_by_todo = {row.todo_id: row for row in fills}
    records = []
    payloads = []
    for todo in todos:
        fill = fill_by_todo.get(todo.todo_id)
        data = (fill.data_json or {}) if fill else (todo.payload or {})
        if data:
            payloads.append(data)
        filled_fields = sum(1 for col in columns if _has_value(data.get(col["key"])))
        records.append({
            "todo_id": todo.todo_id,
            "assignee_id": todo.assignee_id,
            "assignee_name": todo_svc._display_name(db, todo.assignee_id),
            "company": todo_svc.row_company(db, todo, task),  # 看板汇总表的公司列
            "status": todo.status,
            "submitted_at": (
                fill.submitted_at.isoformat() if fill and fill.submitted_at
                else todo.submitted_at.isoformat() if todo.submitted_at else None
            ),
            "filled_fields": filled_fields,
            "field_total": len(columns),
            "data": data,
        })

    fields = []
    for col in columns:
        values = []
        filled_payloads = 0
        for payload in payloads:
            raw = payload.get(col["key"])
            present = [value for value in (raw if isinstance(raw, list) else [raw]) if not _is_empty(value)]
            if present:
                filled_payloads += 1
                values.extend(present)
        number_total = sum(float(v) for v in values if isinstance(v, (int, float)) and not isinstance(v, bool))
        options: dict[str, int] = {}
        if col.get("type") in ("select", "date", "text"):
            for value in values:
                label = str(value)
                options[label] = options.get(label, 0) + 1
        fields.append({
            "key": col["key"], "label": col["label"], "type": col.get("type", "text"),
            "filled": filled_payloads, "total": len(todos), "number_total": number_total,
            "options": dict(sorted(options.items(), key=lambda item: (-item[1], item[0]))[:8]),
        })
    return {
        "task_id": task.task_id, "task_title": task.title, "template_name": tpl.name,
        "total": len(todos), "submitted": len(fills),
        "completed": sum(1 for todo in todos if todo.status in ("completed", "na_closed")),
        "records": records, "fields": fields,
        "tabs": (tpl.parse_result_json or {}).get("tabs", []),
    }


def _is_empty(value: object) -> bool:
    return value is None or value == "" or value == [] or value == {}


def _has_value(value: object) -> bool:
    return any(not _is_empty(item) for item in value) if isinstance(value, list) else not _is_empty(value)


def parse_excels(files: list[tuple[str, bytes]]) -> dict:
    """每个文件一个分页；持久化字段键加分页前缀，保持既有扁平 payload 兼容性。"""
    if not 1 <= len(files) <= 10:
        raise errors.validation("请选择 1 至 10 个 Excel 文件")
    columns, tabs = [], []
    for index, (filename, content) in enumerate(files):
        try:
            parsed = parse_excel(content, filename)
        except errors.AppError as exc:
            raise errors.validation(f"文件「{filename}」：{exc.message}") from exc
        tab_id = f"file_{index + 1}"
        fields = [{**c, "key": f"{tab_id}__{c['key']}"} for c in parsed["columns"]]
        columns.extend(fields)
        tabs.append({
            "id": tab_id, "name": Path(filename).stem,
            "source_file_name": filename, "field_keys": [c["key"] for c in fields],
            "sample_rows": [{f"{tab_id}__{key}": value for key, value in row.items()}
                            for row in parsed.get("sample_rows", [])],
            "presentation": parsed["presentation"], "repeatable": parsed["presentation"] == "grid",
            "title": parsed["title"],
            "metadata": parsed["metadata"],
        })
    return {"sheet_name": "多文件填报", "columns": columns, "tabs": tabs}


def apply_tab_edits(parsed: dict, edited: dict) -> dict:
    """仅接受预览中的字段调整，文件归属及布局以服务端重新解析结果为准。"""
    if not isinstance(edited, dict) or not isinstance(edited.get("columns"), list):
        raise errors.validation("字段调整数据格式错误")
    columns = edited["columns"]
    if not all(isinstance(c, dict) for c in columns):
        raise errors.validation("字段调整数据格式错误")
    _validate_columns(columns, MAX_COLUMNS * len(parsed["tabs"]))
    original = {c["key"]: c for c in parsed["columns"]}
    if any(c["key"] not in original for c in columns):
        raise errors.validation("字段不属于本次上传的文件")
    allowed = {"label", "type", "required", "options"}
    normalized = [{**original[c["key"]], **{k: v for k, v in c.items() if k in allowed}} for c in columns]
    retained = {c["key"] for c in normalized}
    tabs = [{**tab, "field_keys": [k for k in tab["field_keys"] if k in retained]} for tab in parsed["tabs"]]
    return {**parsed, "columns": normalized, "tabs": tabs}
